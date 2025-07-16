import os
import win32com.client
import pythoncom
from elasticsearch import Elasticsearch
from elasticsearch.helpers import bulk
from datetime import datetime
import traceback
import threading
from concurrent.futures import ThreadPoolExecutor
from ldap3 import Server, Connection, ALL, SUBTREE
from dotenv import load_dotenv
import mimetypes

import xlrd
# Additional imports for content extraction
import PyPDF2
import docx
import openpyxl

load_dotenv()

ELASTIC_HOST = os.getenv("ELASTIC_HOST")
ELASTIC_USER = os.getenv("ELASTIC_USER")
ELASTIC_PASSWORD = os.getenv("ELASTIC_PASSWORD")
ACTIVE_ADDRESS = os.getenv("ACTIVE_ADDRESS")
ACTIVE_PASSWORD = os.getenv("ACTIVE_PASSWORD")
ACTIVE_SEARCH_BASE = os.getenv("ACTIVE_SEARCH_BASE")
es = Elasticsearch(
    hosts=[ELASTIC_HOST],
    basic_auth=(ELASTIC_USER, ELASTIC_PASSWORD)
)

server = Server(ACTIVE_ADDRESS, get_info=ALL)
conn = Connection(
    server,
    user='alborz\\ldap.user',
    password=ACTIVE_PASSWORD,
    auto_bind=True
)

page_size = 1000
paged_search = conn.extend.standard.paged_search(
    search_base=ACTIVE_SEARCH_BASE,
    search_filter='(objectClass=user)',
    search_scope=SUBTREE,
    attributes=['sAMAccountName', 'displayName', 'mail', 'manager', 'department'],
    paged_size=page_size,
    generator=True
)

ad_users = {}
for entry in paged_search:
    attr = entry.get('attributes', {})
    email = str(attr.get('mail', '')).lower().strip()
    if email:
        ad_users[email] = {
            "display_name": str(attr.get('displayName', '')),
            "sam_account_name": str(attr.get('sAMAccountName', '')),
            "department": str(attr.get('department', '')),
            "manager": str(attr.get('manager', ''))
        }

BULK_SIZE = 1000
MAX_WORKERS = 4

outlook_lock = threading.Lock()
counter_lock = threading.Lock()
total_emails = 0
indexed_emails = 0

def extract_text_from_file(file_path):
    mime_type, _ = mimetypes.guess_type(file_path)
    text = ""

    try:
        if mime_type == 'application/pdf':
            with open(file_path, 'rb') as f:
                reader = PyPDF2.PdfReader(f)
                text = "\n".join([page.extract_text() or "" for page in reader.pages])

        elif mime_type == 'application/vnd.openxmlformats-officedocument.wordprocessingml.document':
            doc = docx.Document(file_path)
            text = "\n".join([p.text for p in doc.paragraphs])

        elif mime_type == 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet':
            wb = openpyxl.load_workbook(file_path, data_only=True)
            for sheet in wb.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    text += " ".join([str(cell) if cell is not None else '' for cell in row]) + "\n"

        elif mime_type == 'application/vnd.ms-excel':
            wb = xlrd.open_workbook(file_path)
            for sheet in wb.sheets():
                for row_idx in range(sheet.nrows):
                    row_values = sheet.row_values(row_idx)
                    text += " ".join([str(cell) for cell in row_values]) + "\n"

    except Exception as e:
        print(f"[!] Error extracting text from {file_path}: {e}")

    return text.strip()

def extract_cn(distinguished_name):
    if distinguished_name and "CN=" in distinguished_name:
        return distinguished_name.split("CN=")[1].split(",")[0]
    return ""

def save_attachments(message, base_path, user_name, email_id):
    attachments_info = []
    if message.Attachments.Count > 0:
        user_folder = os.path.join(base_path, user_name)
        save_path = os.path.join(user_folder, email_id)
        os.makedirs(save_path, exist_ok=True)
        for i in range(1, message.Attachments.Count + 1):
            attachment = message.Attachments.Item(i)
            filename = attachment.FileName
            file_path = os.path.join(save_path, filename)
            attachment.SaveAsFile(file_path)

            attachment_data = {
                "filename": filename,
                "filepath": file_path,
                "size": os.path.getsize(file_path),
            }

            # اضافه کردن متن ضمیمه اگر قابل استخراج بود
            extracted_text = extract_text_from_file(file_path)
            if extracted_text:
                attachment_data["text"] = extracted_text

            attachments_info.append(attachment_data)
    return attachments_info

def clean_email_field(email_field):
    if email_field:
        return [email.strip("'").strip() for email in email_field.split(';') if email.strip()]
    return []

def read_folder(folder, user_name):
    bulk_actions = []
    local_total = 0
    local_indexed = 0

    def index_bulk():
        nonlocal bulk_actions, local_indexed
        if bulk_actions:
            try:
                result = bulk(es, bulk_actions, stats_only=False)
                print(f"Bulk result: {result}")
                local_indexed += len(bulk_actions)
            except Exception as e:
                print(f"Bulk index failed: {e}")
            finally:
                bulk_actions = []

    try:
        target_folders = ["inbox", "sent items", "deleted items"]
        folder_name = folder.Name.lower()
        if folder_name in target_folders:
            messages = folder.Items
            for message in messages:
                try:
                    if message.Class == 43:
                        subject = message.Subject or ""
                        sender = message.SenderName or ""
                        body = (message.Body or "").strip().replace('\n', '')
                        received = message.ReceivedTime
                        email_o = message.SenderEmailAddress
                        email_o_clean = str(email_o).lower().strip()
                        try:
                            exch_user = message.Sender.GetExchangeUser()
                            if exch_user:
                                email_o = exch_user.PrimarySmtpAddress
                                email_o_clean = str(email_o).lower().strip()
                        except Exception as ex:
                            pass
                        attachments = save_attachments(message, r"D:\\attachments", user_name, message.EntryID)
                        if isinstance(received, datetime):
                            received = received.strftime("%Y-%m-%dT%H:%M:%S%z")

                        email_doc = {
                            "subject": subject,
                            "sender": sender,
                            "body": body,
                            "to": clean_email_field(message.To),
                            "cc": clean_email_field(message.CC),
                            "date": received,
                            "user": user_name,
                            "attachments": attachments,
                            "email": email_o,
                            "folder_name": folder_name,
                        }
                        if email_o_clean in ad_users:
                            user_info = ad_users[email_o_clean]
                            email_doc.update({
                                "display_name": user_info.get("display_name", ""),
                                "sam_account_name": user_info.get("sam_account_name", ""),
                                "department": user_info.get("department", ""),
                                "manager": extract_cn(user_info.get("manager", "")), 
                            })

                        bulk_actions.append({
                            "_index": "email_exchange",
                            "_source": email_doc
                        })
                        local_total += 1
                        if len(bulk_actions) >= BULK_SIZE:
                            index_bulk()
                except Exception as e:
                    print(f"Failed to process message: {e}")

        index_bulk()

        for sub_folder in folder.Folders:
            sub_total, sub_indexed = read_folder(sub_folder, user_name)
            local_total += sub_total
            local_indexed += sub_indexed

    except Exception as e:
        print(f"Error reading folder {folder.Name}: {e}")

    return local_total, local_indexed

def extract_emails_from_pst(pst_path, folder_name):
    global total_emails, indexed_emails
    pythoncom.CoInitialize()
    outlook = None
    root_folder = None

    try:
        with outlook_lock:
            outlook = win32com.client.Dispatch("Outlook.Application").GetNamespace("MAPI")
            outlook.AddStore(pst_path)

            for i in range(outlook.Folders.Count):
                if outlook.Folders.Item(i + 1).Name.lower() in os.path.basename(pst_path).lower():
                    root_folder = outlook.Folders.Item(i + 1)
                    break
            if not root_folder:
                root_folder = outlook.Folders.Item(outlook.Folders.Count)

        file_total, file_indexed = read_folder(root_folder, folder_name)
        print(f"[{pst_path}] Total: {file_total}, Indexed: {file_indexed}")

        with counter_lock:
            total_emails += file_total
            indexed_emails += file_indexed

    except Exception as e:
        print(f"Error processing {pst_path}: {e}")
        traceback.print_exc()

    finally:
        with outlook_lock:
            try:
                if outlook and root_folder:
                    outlook.RemoveStore(root_folder)
            except Exception as e:
                print(f"Failed to remove store for {pst_path}: {e}")

def process_pst_file(pst_path):
    folder_name = os.path.basename(os.path.dirname(pst_path))
    print(f"Processing {pst_path}")
    extract_emails_from_pst(pst_path, folder_name)

base_dir = r"D:\\test2"
pst_files = []
for root, dirs, files in os.walk(base_dir):
    for file in files:
        if file.endswith(".pst"):
            full_path = os.path.join(root, file)
            pst_files.append(full_path)

with ThreadPoolExecutor(max_workers=MAX_WORKERS) as executor:
    executor.map(process_pst_file, pst_files)

print("=" * 40)
print(f"Total Emails Processed: {total_emails}")
print(f"Total Emails Indexed:   {indexed_emails}")
print("=" * 40)
